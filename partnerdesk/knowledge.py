"""Brand knowledge base: a brand's own material → typed, self-contained fragments.

Ported from `legacy/brand_knowledge` onto the project's one model seam (`llm.complete_json`),
text layer only for now — the campus text models read the pages' text; page images for a
vision model are the next step (documents.py in legacy shows how).

Two passes, kept from the first version because a one-pass extractor with a fixed schema
quietly drops whatever doesn't fit it (hand it a product schema and the brand story, the
sales proof and the audience pages vanish):

  pass 1 — scan:    "what kinds of content are in here?"  → sections (open type slugs)
  pass 2 — extract: fragments against what pass 1 found    → one fragment per self-contained fact

Long files go in page batches. Everything a run produces is written in one transaction
(`db.save_distillation`) as DRAFT: a person approves a fragment on the Knowledge page
before any prompt may quote it.
"""

from __future__ import annotations

import hashlib
import re
import uuid
from dataclasses import dataclass, field
from typing import Any

from .db import Database
from .llm import LLM
from .models import BrandDocument, DocumentSection, KnowledgeFragment, Product

PAGES_PER_BATCH = 12
MAX_PAGE_CHARS = 6000
MAX_BATCH_CHARS = 40_000
PSEUDO_PAGE_CHARS = 3500  # plain text / docx have no pages; chunk so fragments can still cite one

# Starting taxonomy handed to the model. Not enforced: the model is told to invent a new
# slug when none fits, and unknown types surface on the Knowledge page to be relabelled.
SEED_TYPES: dict[str, str] = {
    "product": "A single product: what it is, selling points, differentiation, specs.",
    "brand_story": "Brand origin, founder story, mission, values, heritage, culture.",
    "proof": "Evidence of traction: sales figures, growth, market share, awards, press.",
    "audience": "Target customer/market profiles — who this is for and why.",
    "certification": "Certifications, test reports, regulatory or compliance status.",
    "case_study": "A named customer, distributor, or market success story.",
    "market": "Category or market context: size, trends, competitive landscape.",
    "usage": "How to use, sell or merchandise the product: routines, dosage, shelf placement, training points.",
    "other": "Genuinely relevant material that fits none of the above.",
}
_SEED_LIST = "\n".join(f"- {slug}: {desc}" for slug, desc in SEED_TYPES.items())

SCAN_SYSTEM = f"""You are cataloguing a brand's own sales, product and marketing material so a distributor-facing assistant can later answer questions about the brand and its products accurately.

Read every page you are given and report what kinds of content are actually present. Do not extract details yet — this pass only identifies what is here.

These are the content types we already know about:
{_SEED_LIST}

Rules:
- Only report a type if the material genuinely contains it. Do not pad the list.
- If a body of content does not fit any type above, invent a new snake_case slug for it and describe it. This matters more than reusing an existing type: a forced fit loses the content later.
- Record the 1-based page numbers where each type appears.
- Ignore pure decoration, agendas, title slides, and legal boilerplate.

Respond with ONLY a raw JSON object (no markdown fence):
{{"summary": string, "types": [{{"type": string, "label": string, "description": string, "page_numbers": [number]}}]}}"""

EXTRACT_SYSTEM = """You are distilling a brand's sales, product and marketing material into reusable knowledge fragments. A few fragments at a time will later be handed to an assistant talking to the brand's distributors, so it has accurate, specific things to say about the brand and its products.

Write one fragment per distinct, self-contained piece of knowledge.

Rules:
- Each fragment must stand alone. A later reader sees only that fragment, never the source page, so never write "as shown above" or "this product" without naming it.
- Be concrete and keep real specifics: figures, percentages, certifications, ingredient or material names, market names, dates, award names. Vague marketing prose is worthless here — "high quality formulation" tells a reader nothing; "22% niacinamide, third-party tested for purity" does.
- Never invent, extrapolate, or round facts that are not in the material. If a claim is qualified or dated in the source, keep the qualifier.
- Use the type slugs identified in the scan. If a fragment genuinely fits none of them, use a new snake_case slug rather than forcing it into a close-enough one.
- Set product_name only for type 'product', copying the product name verbatim as written in the material.
- Record the 1-based source pages for every fragment so a human can verify it.
- Split a page covering several products into one fragment per product. Merge a single product spread across several pages into one fragment.
- Skip decoration, agendas, title slides, table-of-contents pages, and legal boilerplate.
- Write content in the language of the material.

Respond with ONLY a raw JSON object (no markdown fence):
{"fragments": [{"type": string, "title": string, "content": string, "product_name": string|null, "source_pages": [number], "tags": [string]}]}"""

SCAN_SCHEMA: dict[str, Any] = {
    "name": "brand_document_scan",
    "strict": False,
    "schema": {
        "type": "object",
        "properties": {
            "summary": {"type": "string"},
            "types": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "type": {"type": "string"}, "label": {"type": "string"}, "description": {"type": "string"},
                        "page_numbers": {"type": "array", "items": {"type": "integer"}},
                    },
                    "required": ["type", "label", "description", "page_numbers"],
                },
            },
        },
        "required": ["summary", "types"],
    },
}

EXTRACT_SCHEMA: dict[str, Any] = {
    "name": "brand_knowledge_fragments",
    "strict": False,
    "schema": {
        "type": "object",
        "properties": {
            "fragments": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "type": {"type": "string"}, "title": {"type": "string"}, "content": {"type": "string"},
                        "product_name": {"type": ["string", "null"]},
                        "source_pages": {"type": "array", "items": {"type": "integer"}},
                        "tags": {"type": "array", "items": {"type": "string"}},
                    },
                    "required": ["type", "title", "content", "product_name", "source_pages", "tags"],
                },
            },
        },
        "required": ["fragments"],
    },
}

SUPPORTED_EXTS = (".pdf", ".txt", ".md", ".docx")


# --- the file → pages of text -----------------------------------------------------------

def _chunk(text: str, size: int = PSEUDO_PAGE_CHARS) -> list[str]:
    text = text.strip()
    if not text:
        return []
    out, buf = [], ""
    for para in re.split(r"\n\s*\n", text):
        if buf and len(buf) + len(para) + 2 > size:
            out.append(buf)
            buf = ""
        buf = f"{buf}\n\n{para}" if buf else para
    if buf:
        out.append(buf)
    return out


def load_pages(file_name: str, data: bytes) -> list[tuple[int, str]]:
    """(page_number, text) per page. PDF via PyMuPDF; docx / txt / md are chunked into
    pseudo-pages so a fragment can still point somewhere."""
    name = file_name.lower()
    if name.endswith(".pdf"):
        try:
            import pymupdf  # the `contract` extra
        except ImportError as err:
            raise ValueError('Reading PDFs needs PyMuPDF: pip install -e ".[contract]"') from err
        pages: list[tuple[int, str]] = []
        with pymupdf.open(stream=data, filetype="pdf") as doc:
            for i, page in enumerate(doc, start=1):
                pages.append((i, (page.get_text() or "").strip()))
        return pages
    if name.endswith(".docx"):
        import docx  # python-docx, also in the `contract` extra

        d = docx.Document(__import__("io").BytesIO(data))
        text = "\n\n".join(p.text for p in d.paragraphs if p.text.strip())
        for t in d.tables:
            text += "\n\n" + "\n".join(" | ".join(c.text.strip() for c in row.cells) for row in t.rows)
        return list(enumerate(_chunk(text), start=1))
    if name.endswith((".txt", ".md")):
        text = data.decode("utf-8", errors="replace")
        parts = [p for p in text.split("\f") if p.strip()] if "\f" in text else _chunk(text)
        return list(enumerate(parts, start=1))
    raise ValueError(f"Unsupported document: {file_name} (use {', '.join(SUPPORTED_EXTS)})")


# --- the two passes ------------------------------------------------------------------------

@dataclass
class Distillation:
    summary: str
    sections: list[DocumentSection]
    fragments: list[KnowledgeFragment]
    model: str | None
    pages_processed: int
    skipped: list[str] = field(default_factory=list)


def _slug(s: Any) -> str:
    s = re.sub(r"[^a-z0-9]+", "_", str(s or "").strip().lower()).strip("_")
    return s or "other"


def _ints(v: Any) -> list[int]:
    out: list[int] = []
    for x in v if isinstance(v, list) else []:
        try:
            n = int(x)
        except (TypeError, ValueError):
            continue
        if n > 0 and n not in out:
            out.append(n)
    return sorted(out)


def _batch_text(file_name: str, pages: list[tuple[int, str]]) -> str:
    parts = [f"Source file: {file_name}. Pages {pages[0][0]}–{pages[-1][0]}."]
    used = len(parts[0])
    for n, text in pages:
        block = f"--- Page {n} ---\n{(text or '(no text on this page)')[:MAX_PAGE_CHARS]}"
        if used + len(block) > MAX_BATCH_CHARS:
            parts.append(f"--- Page {n} --- (omitted: batch too long)")
            continue
        parts.append(block)
        used += len(block)
    return "\n\n".join(parts)


def scan_pages(llm: LLM, file_name: str, pages: list[tuple[int, str]]) -> tuple[str, list[DocumentSection]]:
    raw = llm.complete_json(system=SCAN_SYSTEM, messages=[{"role": "user", "content": _batch_text(file_name, pages)}], schema=SCAN_SCHEMA, temperature=0.1, max_tokens=1200)
    sections = [
        DocumentSection(type=_slug(t.get("type")), label=str(t.get("label") or t.get("type") or "").strip() or _slug(t.get("type")),
                        description=str(t.get("description") or "").strip() or None, page_numbers=_ints(t.get("page_numbers")))
        for t in raw.get("types") or [] if isinstance(t, dict)
    ]
    return str(raw.get("summary") or "").strip(), sections


def extract_fragments(llm: LLM, file_name: str, pages: list[tuple[int, str]], summary: str, sections: list[DocumentSection]) -> list[dict[str, Any]]:
    found = "\n".join(f"- {s.type} ({s.label}): {s.description or ''} [pages {', '.join(map(str, s.page_numbers)) or 'unspecified'}]" for s in sections)
    user = (
        f"Overview of this material:\n{summary or '(none)'}\n\nContent types identified in it:\n{found or '- (none identified)'}\n\n"
        f"Now extract the knowledge fragments.\n\n{_batch_text(file_name, pages)}"
    )
    raw = llm.complete_json(system=EXTRACT_SYSTEM, messages=[{"role": "user", "content": user}], schema=EXTRACT_SCHEMA, temperature=0.2, max_tokens=3000)
    out: list[dict[str, Any]] = []
    for f in raw.get("fragments") or []:
        if not isinstance(f, dict):
            continue
        content = str(f.get("content") or "").strip()
        title = str(f.get("title") or "").strip()
        if not content or not title:
            continue
        pn = f.get("product_name")
        out.append({
            "type": _slug(f.get("type")), "title": title[:200], "content": content,
            "product_name": (pn.strip()[:200] if isinstance(pn, str) and pn.strip() else None),
            "source_pages": _ints(f.get("source_pages")),
            "tags": [str(t).strip()[:40] for t in (f.get("tags") if isinstance(f.get("tags"), list) else []) if str(t).strip()][:8],
        })
    return out


def match_product(products: list[Product], product_name: str | None, title: str | None = None) -> Product | None:
    """A fragment about a catalog product carries its id. Exact (case-insensitive) name or
    SKU first; otherwise the catalog SKU or name has to appear inside what the model wrote
    ("Radiance Serum (SR-200)"), longest match winning so "Radiance Serum Plus" never
    resolves to "Radiance Serum"."""
    names = [n.strip().lower() for n in (product_name, title) if n and n.strip()]
    if not names:
        return None
    keyed = [(p, {k.strip().lower() for k in (p.sku, p.product_name_en, p.native_name) if k and k.strip()}) for p in products]
    for p, keys in keyed:
        if any(n in keys for n in names):
            return p
    best: tuple[int, Product] | None = None
    for p, keys in keyed:
        for k in keys:
            if len(k) >= 3 and any(re.search(rf"(?<![a-z0-9]){re.escape(k)}(?![a-z0-9])", n) for n in names) and (best is None or len(k) > best[0]):
                best = (len(k), p)
    return best[1] if best else None


def distill(
    llm: LLM, *, brand_id: str, file_name: str, pages: list[tuple[int, str]], document_id: str | None = None,
    products: list[Product] | None = None, pages_per_batch: int = PAGES_PER_BATCH,
) -> Distillation:
    """Both passes over every page batch, merged. Pure with respect to the database."""
    readable = [(n, t) for n, t in pages if (t or "").strip()]
    skipped = [f"page {n}: no text layer (scanned page — needs a vision model)" for n, t in pages if not (t or "").strip()]
    if not readable:
        raise ValueError(f"No readable text in {file_name!r}.")
    sections: dict[str, DocumentSection] = {}
    fragments: list[KnowledgeFragment] = []
    summaries: list[str] = []
    for i in range(0, len(readable), max(1, pages_per_batch)):
        batch = readable[i : i + pages_per_batch]
        summary, found = scan_pages(llm, file_name, batch)
        if summary:
            summaries.append(summary)
        for s in found:
            if s.type in sections:
                sections[s.type] = sections[s.type].model_copy(update={"page_numbers": sorted(set(sections[s.type].page_numbers) | set(s.page_numbers))})
            else:
                sections[s.type] = s
        for f in extract_fragments(llm, file_name, batch, summary, list(sections.values())):
            product = match_product(products or [], f["product_name"], f["title"] if f["type"] == "product" else None)
            fragments.append(KnowledgeFragment(
                id=f"kf_{uuid.uuid4().hex[:12]}", brand_id=brand_id, document_id=document_id, product_id=product.id if product else None, **f,
            ))
    model = getattr(getattr(llm, "settings", None), "model", None)
    return Distillation(summary=" ".join(summaries).strip(), sections=list(sections.values()), fragments=fragments, model=model, pages_processed=len(readable), skipped=skipped)


# --- database ------------------------------------------------------------------------------

def ingest_document(db: Database, *, brand_id: str, file_name: str, data: bytes, title: str | None = None, media_type: str | None = None) -> BrandDocument:
    """Store the upload as a document row plus its text per page. Distillation is a separate,
    explicit step (it costs model calls; a person presses the button)."""
    pages = load_pages(file_name, data)
    doc = BrandDocument(
        id=f"doc_{uuid.uuid4().hex[:12]}", brand_id=brand_id, title=title or file_name, file_name=file_name,
        file_sha256=hashlib.sha256(data).hexdigest(), media_type=media_type, size_bytes=len(data), page_count=len(pages),
    )
    db.save_document(doc, pages)
    return doc


def distill_document(db: Database, llm: LLM, document_id: str, *, products: list[Product] | None = None) -> Distillation:
    doc = db.document(document_id)
    if not doc:
        raise KeyError(f"document {document_id} not found")
    pages = db.document_pages(document_id)
    try:
        result = distill(llm, brand_id=doc.brand_id, file_name=doc.file_name, pages=pages, document_id=document_id, products=products)
    except Exception:
        db.save_distillation(document_id, summary=doc.summary, model=doc.model, pages_processed=doc.pages_processed, sections=[], fragments=[], status="failed")
        raise
    db.save_distillation(document_id, summary=result.summary or None, model=result.model, pages_processed=result.pages_processed, sections=result.sections, fragments=result.fragments)
    return result


def add_fragment(
    db: Database, *, brand_id: str, type: str, title: str, content: str, product_name: str | None = None,
    tags: list[str] | None = None, products: list[Product] | None = None, status: str = "approved",
) -> KnowledgeFragment:
    """A fragment a person typed in. Theirs, so approved by default."""
    product = match_product(products or [], product_name, title)
    f = KnowledgeFragment(
        id=f"kf_{uuid.uuid4().hex[:12]}", brand_id=brand_id, type=_slug(type), title=title.strip()[:200], content=content.strip(),
        product_name=product_name, product_id=product.id if product else None, tags=tags or [], status=status,  # type: ignore[arg-type]
    )
    db.upsert_fragment(f)
    return f
