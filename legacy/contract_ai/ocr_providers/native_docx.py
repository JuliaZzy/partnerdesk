"""DOCX → NormalizedDocument via structured text extraction (no OCR)."""

from __future__ import annotations

from pathlib import Path

from ..schemas import (
    BlockType,
    DerivedArtifacts,
    DocumentMetadata,
    NormalizedDocument,
    NormalizedPage,
    TextBlock,
    utc_now_iso,
)


def docx_to_normalized_document(file_path: str) -> NormalizedDocument:
    """
    Extract plain text from a .docx file (Office Open XML).
    Legacy .doc is not supported here — convert to .docx or PDF.
    """
    try:
        from docx import Document as DocxDocument
    except ImportError as e:
        raise RuntimeError("Install python-docx for DOCX support: pip install python-docx") from e

    path = Path(file_path)
    doc = DocxDocument(str(path))

    lines: list[str] = []
    for para in doc.paragraphs:
        t = (para.text or "").strip()
        if t:
            lines.append(t)

    for table in doc.tables:
        for row in table.rows:
            cells = [(c.text or "").strip() for c in row.cells]
            row_text = "\t".join(c for c in cells if c)
            if row_text:
                lines.append(row_text)

    full = "\n".join(lines).strip()
    page_text = full if full else ""
    blocks: list[TextBlock] = (
        [TextBlock(text=page_text, block_type=BlockType.PARAGRAPH)] if page_text else []
    )

    return NormalizedDocument(
        source_kind="native_text",
        ocr_provider_id="native_docx",
        ocr_provider_version="python-docx",
        source_filename=path.name,
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        full_text=page_text,
        pages=[
            NormalizedPage(
                page_number=1,
                text=page_text,
                blocks=blocks,
            )
        ],
        derived=DerivedArtifacts(),
        metadata=DocumentMetadata.model_validate(
            {
                "native_docx": True,
                "processed_at": utc_now_iso(),
            }
        ),
    )
